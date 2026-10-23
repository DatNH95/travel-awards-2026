"""Read the supplied workbook without modifying it; generate category form data."""
import json
import re
import sys
from pathlib import Path
import openpyxl

workbook = openpyxl.load_workbook(sys.argv[1], data_only=True)
keys = [
    'Điểm đến du lịch của năm', 'Doanh nghiệp lữ hành của năm', 'Khách sạn của năm',
    'Khu nghỉ dưỡng của năm (Resort)', 'Hãng hàng không du lịch của năm', 'Trải nghiệm du lịch của năm',
    'Nhân vật du lịch của năm (Cá nhân)', 'Sáng kiến du lịch bền vững của năm',
    'Điểm đến du lịch xanh của năm', 'Sản phẩm du lịch sáng tạo của năm', 'Chuyển đổi số xuất sắc trong du lịch',
    'Làng du lịch cộng đồng tiêu biểu của năm', 'Trải nghiệm ẩm thực du lịch của năm',
    'Gương mặt truyền thông du lịch của năm (Đại sứ)', 'Chiến dịch quảng bá du lịch của năm',
]
def bullets(value):
    items = []
    for line in (value or '').splitlines():
        line = line.strip()
        if not line:
            continue
        if line.startswith('•') or not items:
            items.append(line.lstrip('•').strip())
        else:
            items[-1] += ' ' + line
    return [clean(item) for item in items]

def clean(value):
    for old, new in [('phân khúc) )', 'phân khúc)'), ('&', 'và'), ('TripAdvisor', 'Tripadvisor'), ('nhà hàng/ điểm đến', 'nhà hàng/điểm đến'), ('≥ 1600px', 'có kích thước tối thiểu 1.600 px'), ('≤ 3 phút', 'không quá 3 phút'), ('MỘT', 'một'), (' + ', ', '), ('(baseline)', '(mốc ban đầu)'), ('nền tảng review', 'nền tảng đánh giá'), ('nguồn Cục Hàng không, OAG/Cirium hoặc nội bộ', 'nguồn: Cục Hàng không, OAG/Cirium hoặc báo cáo nội bộ')]:
        value = value.replace(old, new)
    return value.strip()

records = {}
labels = [
    ['Số liệu khách và tổng thu du lịch', 'Sản phẩm, hạ tầng và kết nối giao thông mới', 'Chính sách và chương trình xúc tiến', 'Hoạt động bảo tồn và phát triển cộng đồng', 'Kênh truyền thông chính thức'],
    ['Doanh thu và lượt khách', 'Kết quả theo phân khúc', 'Sản phẩm tour tiêu biểu', 'Quy mô nhân sự', 'Kênh truyền thông chính thức'],
    ['Quy mô và hiệu quả kinh doanh', 'Đánh giá trên các nền tảng OTA', 'Quy mô và đào tạo nhân sự', 'Kênh truyền thông chính thức'],
    ['Quy mô và hiệu quả kinh doanh', 'Đánh giá trên các nền tảng OTA', 'Quy mô và đào tạo nhân sự', 'Kênh truyền thông chính thức'],
    ['Mạng đường bay và tần suất', 'Số hành khách vận chuyển', 'Tỷ lệ chuyến bay đúng giờ (OTP)', 'Chương trình giá và kích cầu', 'Hợp tác quảng bá điểm đến', 'Kênh truyền thông chính thức'],
    ['Lượng khách và tăng trưởng', 'Hiệu quả truyền thông', 'Đánh giá trên các nền tảng', 'Kênh truyền thông chính thức'],
    [],
    ['Sáng kiến và phạm vi triển khai', 'Kết quả và tác động', 'So sánh trước và sau triển khai'],
    ['Chứng nhận xanh hoặc bền vững', 'Kết quả phát triển du lịch xanh'],
    ['Giới thiệu sản phẩm', 'Hiệu quả kinh doanh', 'Kênh truyền thông chính thức'],
    ['Giải pháp và thời gian triển khai', 'Công nghệ sử dụng', 'Hiệu quả vận hành'],
    ['Giới thiệu mô hình', 'Kết quả hoạt động du lịch', 'Sự tham gia và sinh kế cộng đồng'],
    ['Doanh thu và lượt khách', 'Mô hình và trải nghiệm ẩm thực', 'Giải thưởng hoặc xếp hạng', 'Kênh truyền thông chính thức', 'Đánh giá trên các nền tảng'],
    [],
    ['Chiến dịch và thời gian triển khai', 'KPI mục tiêu và kết quả', 'Hiệu quả truyền thông và kích cầu'],
]
rows = list(workbook['Giải trụ cột'].iter_rows(min_row=4, values_only=True)) + list(workbook['Giải tiên phong'].iter_rows(min_row=4, values_only=True))
for key, row in zip(keys, rows, strict=True):
    title, conditions, declarations, criteria, evidence = row[:5]
    if key in (keys[6], keys[13]):
        continue
    narrative_labels = [clean(re.sub(r'^\d+\.\s*', '', line.strip())) for line in criteria.splitlines() if re.match(r'^\d+\.', line.strip())]
    records[key] = {
        'audience': clean(title.split('Đối tượng:', 1)[1]),
        'conditions': bullets(conditions),
        'fields': [{'label': label, 'type': 'textarea', 'help': help} for label, help in zip(labels[keys.index(key)], bullets(declarations), strict=True)],
        'narratives': [{'label': label, 'type': 'textarea', 'maxWords': 300} for label in narrative_labels],
        'evidence': bullets(evidence),
    }
    # Editorial copy for the 22.10 changes; keep declarations in their existing fields.
    for field in records[key]['fields']:
        if 'số lượng follow' in field['help']:
            subject = 'điểm đến' if key == keys[0] else 'đơn vị'
            field['help'] = f'Cung cấp link website, fanpage và tài khoản mạng xã hội chính thức của {subject}; ghi rõ số người theo dõi của từng tài khoản mạng xã hội.'
    if key == keys[0]:
        records[key]['conditions'][2] = 'Không có sự cố an ninh, an toàn du khách đặc biệt nghiêm trọng trong kỳ xét giải (01/10/2025–30/09/2026). Sự cố nghiêm trọng được hiểu là có thương vong hoặc có văn bản chỉ đạo xử lý của cơ quan Trung ương (Bộ VHTTDL, Bộ Công an, Cục Du lịch Quốc gia Việt Nam).'
        records[key]['fields'][3]['help'] = 'Nêu tối đa 5 hoạt động, chương trình bảo tồn di sản, môi trường và phát triển cộng đồng tiêu biểu. Cung cấp báo cáo kết quả hoặc số liệu minh chứng tại mục Hồ sơ & minh chứng.'
    if key == keys[4]:
        records[key]['fields'][4]['help'] = 'Nêu các hoạt động hợp tác quảng bá điểm đến với địa phương, doanh nghiệp. Cung cấp báo cáo giới thiệu thỏa thuận hợp tác với các điểm đến tại mục Hồ sơ & minh chứng.'
    if key == keys[5]:
        records[key]['fields'][1]['help'] = 'Cung cấp số liệu lan tỏa trên truyền thông, mạng xã hội trong kỳ xét giải. Nếu có báo cáo truyền thông, dán link tại mục Hồ sơ & minh chứng.'
    # User override: all categories require 3–10 images, regardless of sheet counts.
    records[key]['images'] = {'min': 3, 'max': 10}
    records[key]['evidence'] = [re.sub(r'Tối thiểu (?:12|15|8) ảnh', '3–10 ảnh', item).replace('Bộ ảnh điểm đến (khuyến khích): 8–15 ảnh', 'Bộ ảnh điểm đến (bắt buộc): 3–10 ảnh').replace('tối đa 15 ảnh', 'tối đa 10 ảnh') for item in records[key]['evidence']]
    records[key]['menuRequired'] = key == 'Trải nghiệm ẩm thực du lịch của năm'

out = Path('src/app/(phase-1)/dang-ky-de-cu/category-criteria.ts')
out.write_text('''// Source: Travel Awards - Bộ tiêu chí update 22.10.xlsx, Giải trụ cột!A4:E9 / Giải tiên phong!A4:E12.
// Verified 2026-10-23, including red text; regenerate with scripts/read-registration-criteria.py <workbook>.
import type { RegistrationField } from './registration-data';
export type CategoryCriteria = {
  audience: string; conditions: string[]; fields: RegistrationField[];
  narratives: RegistrationField[]; evidence: string[];
  images: { min: number; max: number | null }; menuRequired: boolean;
};
export const categoryCriteria: Record<string, CategoryCriteria> = ''' + json.dumps(records, ensure_ascii=False, indent=2) + ';\n', encoding='utf-8')
